"""Excel da "Foto dia a dia" na Evolução do estoque (pedido do Gabriel, 06/10/2026).

O que trava:
- mesma restrição da aba: não-admin com a área de Estoque toma 403 (o dado histórico não pode
  vazar por URL direta);
- o arquivo é a MESMA série da tela (mesma função, mesmo recorte) e as células são NÚMEROS — um
  "R$ 4.774.667,53" em texto não soma no Excel;
- o recorte da tela vai no cabeçalho do arquivo (unidade e filtros), senão quem recebe o Excel
  não sabe que é só a curva A de um comprador;
- a tela tem o botão, e ele usa a querystring da própria série (`S.evoQS`).
"""
import io
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tests.test_evolucao_acesso import admin_compras, viewer_compras  # noqa: F401 (fixtures)


def test_nao_admin_nao_baixa(client, viewer_compras):  # noqa: F811
    assert client.get("/estoque/api/evolucao.xlsx").status_code == 403


def test_admin_baixa_a_mesma_serie_da_tela_com_numeros(client, admin_compras):  # noqa: F811
    tela = client.get("/estoque/api/evolucao").get_json()
    r = client.get("/estoque/api/evolucao.xlsx")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["Content-Type"]
    assert "evolucao_estoque" in r.headers["Content-Disposition"]
    ws = load_workbook(io.BytesIO(r.data)).active
    linhas = list(ws.iter_rows(values_only=True))
    cab_i = next(i for i, ln in enumerate(linhas) if ln and ln[0] == "Data")
    cab = linhas[cab_i]
    for col in ("Estoque R$", "Parado R$", "% ruptura", "Rup. A %", "Rup. s/ prov.", "Desacel. R$",
                "Ocupação %", "A vencer R$", "Vencido no dia R$", "Pedidos abertos R$"):
        assert col in cab, col
    dados = [ln for ln in linhas[cab_i + 1:] if ln and ln[0]]
    assert len(dados) == len(tela["dias"])                     # um dia por linha, nada a mais
    if dados:
        i = cab.index("Estoque R$")
        assert isinstance(dados[0][i], (int, float))           # número, não texto formatado
        assert str(dados[0][0]) == sorted((d["data"] for d in tela["dias"]), reverse=True)[0]


def test_recorte_vai_no_cabecalho(client, admin_compras):  # noqa: F811
    r = client.get("/estoque/api/evolucao.xlsx?curva=A&comprador_cod=47")
    ws = load_workbook(io.BytesIO(r.data)).active
    topo = " ".join(str(c) for ln in list(ws.iter_rows(values_only=True))[:4] for c in ln if c)
    assert "curva A" in topo and "47" in topo


def test_tela_tem_o_botao():
    js = Path("static/estoque/estoque.js").read_text(encoding="utf-8")
    assert "/estoque/api/evolucao.xlsx?" in js and "S.evoQS" in js
